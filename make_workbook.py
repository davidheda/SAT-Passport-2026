#!/usr/bin/env python3
"""Creates passport.xlsx (run once; afterwards edit the workbook by hand)."""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

COMPS = [
 ("A1","Linear equations in one variable","S1","Algebra","35 %"),
 ("A2","Linear equations in two variables","S2","Algebra","35 %"),
 ("A3","Linear functions","S2","Algebra","35 %"),
 ("A4","Systems of two linear equations","S3","Algebra","35 %"),
 ("A5","Linear inequalities","S3","Algebra","35 %"),
 ("M1","Equivalent expressions","S4","Advanced Math","35 %"),
 ("M2","Nonlinear equations and systems","S5","Advanced Math","35 %"),
 ("M3","Nonlinear functions","S6","Advanced Math","35 %"),
 ("D1","Ratios, rates, proportions, units","S7","Problem-Solving & Data Analysis","15 %"),
 ("D2","Percentages","S7","Problem-Solving & Data Analysis","15 %"),
 ("D3","One-variable data","S8","Problem-Solving & Data Analysis","15 %"),
 ("D4","Two-variable data, scatterplots","S8","Problem-Solving & Data Analysis","15 %"),
 ("D5","Probability, conditional probability","S9","Problem-Solving & Data Analysis","15 %"),
 ("D6","Inference, margin of error","S9","Problem-Solving & Data Analysis","15 %"),
 ("D7","Evaluating statistical claims","S9","Problem-Solving & Data Analysis","15 %"),
 ("G1","Area and volume","S10","Geometry & Trigonometry","15 %"),
 ("G2","Lines, angles, triangles","S10","Geometry & Trigonometry","15 %"),
 ("G3","Right triangles and trigonometry","S10","Geometry & Trigonometry","15 %"),
 ("G4","Circles","S10","Geometry & Trigonometry","15 %"),
]
BLUE = PatternFill("solid", fgColor="1F4E79"); GREY = PatternFill("solid", fgColor="F2F2F2")
WHITE = Font(bold=True, color="FFFFFF"); BOLD = Font(bold=True)
thin = Side(style="thin", color="CCCCCC"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = openpyxl.Workbook()

# ---- README
ws = wb.active; ws.title = "README"
lines = [
 "SAT Math — competence passport workbook (one workbook per class)",
 "",
 "Sheet Passport : one row per student. Enter the SCORE OUT OF 10 of each validation test:",
 "   <code> T1 = first attempt, <code> R = retake. Leave blank if not taken. The level 0/1/2 is",
 "   computed by build_passport.py (9-10 -> 2, 6-8 -> 1, 0-5 -> 0, best of the two attempts).",
 "   Slug = unguessable address of the student's page; leave EMPTY, the script fills it once.",
 "   Diagnostic / Mock 1 / Mock 2 / Target : free text (e.g. 14/22, 29/44, 650).",
 "Sheet Competences : the 19 codes (edit names/sessions here, they feed the pages).",
 "Sheet Config : base_url = address of the GitHub Pages site (must end with /), course_title, teacher.",
 "",
 "Update the site :   ./update.sh        (= python3 build_passport.py + git add/commit/push)",
 "QR codes to print : qr/qr-sheet.html   (one per student, stick in the exercise book)",
 "Teacher overview  : overview.html      (local file, never published: the site root shows nothing)",
]
for i, l in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=l).font = BOLD if i == 1 else Font()
ws.column_dimensions["A"].width = 110

# ---- Config
ws = wb.create_sheet("Config")
ws.append(["key", "value"])
ws.append(["base_url", "https://davidheda.github.io/SAT-Passport-2026/"])
ws.append(["course_title", "SAT Math — Competence Passport"])
ws.append(["teacher", "N. Hedayati — Ecolint LGB"])
for c in ws[1]: c.font = WHITE; c.fill = BLUE
ws.column_dimensions["A"].width = 16; ws.column_dimensions["B"].width = 60

# ---- Competences
ws = wb.create_sheet("Competences")
ws.append(["Code", "Competence", "Session", "Domain", "Weight"])
for c in COMPS: ws.append(list(c))
for c in ws[1]: c.font = WHITE; c.fill = BLUE
for col, w in zip("ABCDE", (8, 40, 9, 34, 9)): ws.column_dimensions[col].width = w

# ---- Passport
ws = wb.create_sheet("Passport", 0)
header = ["Slug", "Name", "Class", "Diagnostic", "Mock 1", "Mock 2", "Target"]
for code, *_ in COMPS: header += [f"{code} T1", f"{code} R"]
ws.append(header)
# sample students (delete them)
samples = [
 ["", "Sample Student A", "Y11 SAT", "14/22", "", "", "600", 7, "", ] ,
 ["", "Sample Student B", "Y11 SAT", "18/22", "", "", "700", 10, "", ],
]
for s in samples: ws.append(s + [""] * (len(header) - len(s)))
for c in ws[1]: c.font = WHITE; c.fill = BLUE; c.alignment = Alignment(horizontal="center", wrap_text=True)
ws.row_dimensions[1].height = 30
ws.freeze_panes = "C2"
ws.column_dimensions["A"].width = 10; ws.column_dimensions["B"].width = 26; ws.column_dimensions["C"].width = 9
for i in range(4, 8): ws.column_dimensions[get_column_letter(i)].width = 11
for i in range(8, len(header) + 1): ws.column_dimensions[get_column_letter(i)].width = 6.5
# score cells: whole number 0..10, colour-coded like the passport
dv = DataValidation(type="whole", operator="between", formula1="0", formula2="10", allow_blank=True)
ws.add_data_validation(dv)
rng = f"H2:{get_column_letter(len(header))}200"
dv.add(rng)
ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["9"], fill=PatternFill("solid", fgColor="B7E1B1")))
ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["6", "8"], fill=PatternFill("solid", fgColor="F8D7A8")))
ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["0", "5"], fill=PatternFill("solid", fgColor="F5B7B1")))
# grey band on retake columns
for i in range(9, len(header) + 1, 2):
    ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor="4F6F9F")

wb.save("passport.xlsx")
print("passport.xlsx written")
