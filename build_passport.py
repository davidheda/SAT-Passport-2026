#!/usr/bin/env python3
"""
SAT Math — competence passport builder.

Reads passport.xlsx (one row per student, two columns per competence code)
and writes:
  overview.html              teacher overview (class grid) — local only, NOT published
  docs/<slug>/index.html     one page per student, unguessable address
  qr/<slug>.png              QR code of each student's page
  qr/qr-sheet.html           printable sheet of all QR codes (stick in exercise books)

Usage:  python3 build_passport.py            (uses passport.xlsx in this folder)
        python3 build_passport.py other.xlsx

Requires: openpyxl, segno   (pip install openpyxl segno)
"""

import sys, os, re, secrets, datetime, html
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
import segno

HERE = Path(__file__).resolve().parent
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "passport.xlsx"
SITE = HERE / "docs"   # GitHub Pages serves /docs (or root) only
QR = HERE / "qr"

# ------------------------------------------------------------------ config
def read_config(wb):
    cfg = {}
    for row in wb["Config"].iter_rows(min_row=2, values_only=True):
        if row[0]:
            cfg[str(row[0]).strip()] = "" if row[1] is None else str(row[1]).strip()
    cfg.setdefault("base_url", "https://example.github.io/sat-passport/")
    cfg.setdefault("course_title", "SAT Math — Competence Passport")
    cfg.setdefault("teacher", "")
    if not cfg["base_url"].endswith("/"):
        cfg["base_url"] += "/"
    return cfg

def read_competences(wb):
    comps = []
    for row in wb["Competences"].iter_rows(min_row=2, values_only=True):
        if row[0]:
            comps.append(dict(code=str(row[0]).strip(), name=row[1], session=row[2],
                              domain=row[3], weight=row[4]))
    return comps

def level(score):
    """0/1/2 from a score out of 10; None if blank."""
    if score is None or score == "":
        return None
    s = float(score)
    return 2 if s >= 9 else 1 if s >= 6 else 0

# ------------------------------------------------------------------ students
def read_students(wb, comps):
    ws = wb["Passport"]
    header = [c.value for c in ws[1]]
    col = {str(h).strip(): i for i, h in enumerate(header) if h is not None}
    students = []
    changed = False
    for r, row in enumerate(ws.iter_rows(min_row=2), start=2):
        vals = [c.value for c in row]
        name = vals[col["Name"]]
        if not name:
            continue
        slug = vals[col["Slug"]]
        if not slug:  # unguessable, generated once and written back
            slug = secrets.token_hex(4)
            ws.cell(row=r, column=col["Slug"] + 1, value=slug)
            changed = True
        st = dict(name=str(name).strip(), slug=str(slug).strip(),
                  cls=vals[col.get("Class", -1)] if "Class" in col else "",
                  diag=vals[col.get("Diagnostic", -1)] if "Diagnostic" in col else None,
                  mock1=vals[col.get("Mock 1", -1)] if "Mock 1" in col else None,
                  mock2=vals[col.get("Mock 2", -1)] if "Mock 2" in col else None,
                  target=vals[col.get("Target", -1)] if "Target" in col else None,
                  comps={})
        for c in comps:
            t1 = vals[col[f"{c['code']} T1"]] if f"{c['code']} T1" in col else None
            rt = vals[col[f"{c['code']} R"]] if f"{c['code']} R" in col else None
            best = None
            for s in (t1, rt):
                if s not in (None, ""):
                    best = s if best is None else max(float(best), float(s))
            st["comps"][c["code"]] = dict(t1=t1, r=rt, best=best, level=level(best))
        students.append(st)
    return students, changed

# ------------------------------------------------------------------ html
CSS = """
:root{--blue:#1F4E79;--gold:#C58B00;--g0:#f5b7b1;--g1:#f8d7a8;--g2:#b7e1b1;--grey:#f2f2f2;--ink:#222}
*{box-sizing:border-box}body{font-family:-apple-system,Helvetica,Arial,sans-serif;margin:0;background:#fafafa;color:var(--ink)}
.wrap{max-width:760px;margin:0 auto;padding:16px}
h1{color:var(--blue);font-size:1.5rem;margin:0 0 2px}.sub{color:#666;font-size:.9rem;margin:0 0 12px}
.legend span{display:inline-block;padding:2px 8px;border-radius:4px;margin-right:6px;font-size:.85rem;font-weight:600}
.l0{background:var(--g0)}.l1{background:var(--g1)}.l2{background:var(--g2)}.lx{background:#fff;border:1px solid #ccc;color:#999}
table{width:100%;border-collapse:collapse;margin:12px 0;font-size:.95rem}
th,td{padding:7px 8px;border-bottom:1px solid #e3e3e3;text-align:left}th{background:var(--grey);font-size:.8rem;text-transform:uppercase;letter-spacing:.03em}
tr.dom td{background:#e6edf5;color:var(--blue);font-weight:700}
td.c{text-align:center;white-space:nowrap}td.lvl{text-align:center;font-weight:700;width:3.2em}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px;margin:10px 0}
.card{background:#fff;border:1px solid #e3e3e3;border-radius:8px;padding:10px}.card b{display:block;font-size:.75rem;color:#666;text-transform:uppercase}
.card span{font-size:1.3rem;font-weight:700;color:var(--blue)}
.prio{background:#fff7e0;border-left:4px solid var(--gold);padding:10px 12px;border-radius:0 6px 6px 0}
.prio ol{margin:6px 0 0 18px;padding:0}
.bar{height:10px;background:#e3e3e3;border-radius:5px;overflow:hidden;margin-top:6px}.bar i{display:block;height:100%;background:var(--blue)}
footer{color:#888;font-size:.8rem;margin-top:20px}
.wheel{width:100%;max-width:600px;display:block;margin:6px auto 0}
.wheel text{font-family:-apple-system,Helvetica,Arial,sans-serif}
.wheel .wc,.wheel .wl,.wheel .cn,.wheel .cs{text-anchor:middle}
.wheel .wc{font-size:17px;font-weight:700;fill:#222}.wheel .wc.s{font-size:12px}
.wheel .wl{font-size:14px;fill:#333}.wheel .dl{font-size:17px;font-weight:700}.wheel .dw{font-size:12px;fill:#777}
.wheel .cn{font-size:40px;font-weight:800;fill:#1F4E79}.wheel .cd{font-size:20px;font-weight:600;fill:#555}.wheel .cs{font-size:11px;letter-spacing:.12em;fill:#777}
.wheel g.w{cursor:pointer}.wheel g.w path{stroke:#fff;stroke-width:1.5;transition:opacity .15s}.wheel g.w:hover path{opacity:.8}
.tip{text-align:center;min-height:1.4em;font-weight:600;color:#1F4E79;margin:4px 0 0}
.hint{text-align:center;color:#666;font-size:.9rem;margin:2px 0 14px}
@media print{body{background:#fff}.wrap{max-width:none}}
"""

def esc(x):
    return html.escape("" if x is None else str(x))

def fmt_score(s):
    if s in (None, ""):
        return ""
    s = float(s)
    return f"{int(s) if s == int(s) else s}/10"


# ------------------------------------------------------------------ wheel (compass)
import math
SHORT = {"Algebra": "Algebra", "Advanced Math": "Advanced Math",
         "Problem-Solving & Data Analysis": "Data Analysis", "Geometry & Trigonometry": "Geometry & Trig"}
RING = {"Algebra": "#1F4E79", "Advanced Math": "#7B3F99",
        "Problem-Solving & Data Analysis": "#C58B00", "Geometry & Trigonometry": "#2E7D32"}
FILL = {None: "#e9e9e9", 0: "#f5b7b1", 1: "#f8d7a8", 2: "#b7e1b1"}

def _pt(r, a):
    a = math.radians(a - 90)
    return r * math.cos(a), r * math.sin(a)

def _arc(r0, r1, a0, a1):
    large = 1 if a1 - a0 > 180 else 0
    x0, y0 = _pt(r1, a0); x1, y1 = _pt(r1, a1); x2, y2 = _pt(r0, a1); x3, y3 = _pt(r0, a0)
    return (f"M{x0:.1f},{y0:.1f} A{r1},{r1} 0 {large} 1 {x1:.1f},{y1:.1f} "
            f"L{x2:.1f},{y2:.1f} A{r0},{r0} 0 {large} 0 {x3:.1f},{y3:.1f}Z")

def wheel_svg(st, comps):
    doms = []
    for c in comps:
        if not doms or doms[-1][0] != c["domain"]:
            doms.append((c["domain"], c["weight"], []))
        doms[-1][2].append(c)
    w = {d: float(str(wt).replace("%", "").strip() or 25) for d, wt, _ in doms}
    tot = sum(w.values())
    R0, R1, RR = 78, 190, 200
    out = []; labels = []; a = 0.0
    n_val = sum(1 for c in comps if st["comps"][c["code"]]["level"] == 2)
    for dom, wt, cs in doms:
        step = 360 / len(comps)          # every competence gets the same wedge
        span = step * len(cs)
        # domain ring
        out.append(f'<path d="{_arc(RR, RR + 9, a + .6, a + span - .6)}" fill="{RING.get(dom, "#555")}"/>')
        for i, c in enumerate(cs):
            b0, b1 = a + i * step + .7, a + (i + 1) * step - .7
            d = st["comps"][c["code"]]; lv = d["level"]
            tip = f'{c["code"]} · {c["name"]} — ' + ("not yet covered" if lv is None else
                  f'level {lv} ({fmt_score(d["best"])})')
            mid = (b0 + b1) / 2
            narrow = step < 12
            rt = (R0 + R1) / 2 + (18 if narrow else 0)
            tx, ty = _pt(rt, mid)
            if narrow:   # radial text for thin wedges
                rot = mid - 90 if mid < 180 else mid + 90
                txt = (f'<text x="{tx:.1f}" y="{ty:.1f}" transform="rotate({rot:.1f} {tx:.1f} {ty:.1f})" '
                       f'class="wc s">{esc(c["code"])}{"" if lv is None else " · " + str(lv)}</text>')
            else:
                txt = (f'<text x="{tx:.1f}" y="{ty - 4:.1f}" class="wc">{esc(c["code"])}</text>'
                       + ("" if lv is None else f'<text x="{tx:.1f}" y="{ty + 14:.1f}" class="wl">{lv}</text>'))
            out.append(f'<g class="w" data-tip="{esc(tip)}"><title>{esc(tip)}</title>'
                       f'<path d="{_arc(R0, R1, b0, b1)}" fill="{FILL[lv]}"/>{txt}</g>')
        # domain label outside the ring
        m = a + span / 2
        lx, ly = _pt(RR + 22, m)
        sx = math.sin(math.radians(m))
        anchor = "middle" if abs(sx) < .3 else ("start" if sx > 0 else "end")
        dy = -12 if ly < -150 else (16 if ly > 150 else 0)
        labels.append(f'<text x="{lx:.1f}" y="{ly + dy:.1f}" text-anchor="{anchor}" class="dl" fill="{RING.get(dom, "#555")}">{esc(SHORT.get(dom, dom))}</text>'
                      f'<text x="{lx:.1f}" y="{ly + dy + 15:.1f}" text-anchor="{anchor}" class="dw">{esc(wt)} of the test</text>')
        a += span
    center = (f'<circle r="{R0 - 6}" fill="#fff"/><text y="6" class="cn">{n_val}<tspan class="cd"> / {len(comps)}</tspan></text>'
              f'<text y="30" class="cs">VALIDATED</text>')
    return (f'<svg class="wheel" viewBox="-318 -262 636 530" role="img" aria-label="Competence wheel">'
            + "".join(out) + center + "".join(labels) + "</svg>")

def student_page(st, comps, cfg, today):
    domains = []
    for c in comps:
        if not domains or domains[-1][0] != c["domain"]:
            domains.append((c["domain"], c["weight"], []))
        domains[-1][2].append(c)
    rows = []
    n_val = n_done = 0
    for dom, weight, cs in domains:
        rows.append(f'<tr class="dom"><td colspan="6">{esc(dom)} — {esc(weight)}</td></tr>')
        for c in cs:
            d = st["comps"][c["code"]]
            lv = d["level"]
            if lv is not None:
                n_done += 1
                n_val += lv == 2
            cls = f"l{lv}" if lv is not None else "lx"
            rows.append(
                f'<tr><td><b>{esc(c["code"])}</b></td><td>{esc(c["name"])}</td><td class="c">{esc(c["session"])}</td>'
                f'<td class="c">{fmt_score(d["t1"])}</td><td class="c">{fmt_score(d["r"])}</td>'
                f'<td class="lvl {cls}">{"" if lv is None else lv}</td></tr>')
    # priority list: level 0 first, then 1; heavier domains first; keep syllabus order inside
    order = {"35 %": 0, "35%": 0, "15 %": 1, "15%": 1}
    prio = [c for c in comps if st["comps"][c["code"]]["level"] in (0, 1)]
    prio.sort(key=lambda c: (st["comps"][c["code"]]["level"], order.get(str(c["weight"]).strip(), 2)))
    prio_html = ("<ol>" + "".join(f"<li><b>{esc(c['code'])}</b> {esc(c['name'])} "
                 f"(level {st['comps'][c['code']]['level']})</li>" for c in prio) + "</ol>") if prio else "<p>Nothing to retake for now.</p>"
    pct = round(100 * n_val / len(comps))
    page = f"""<!DOCTYPE html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><meta name="robots" content="noindex,nofollow">
<title>{esc(st['name'])} — {esc(cfg['course_title'])}</title><style>{CSS}</style></head><body><div class="wrap">
<h1>{esc(cfg['course_title'])}</h1>
<p class="sub"><b>{esc(st['name'])}</b> · {esc(st['cls'])} · updated {today}</p>
<p class="legend"><span class="l2">2 validated (9–10/10)</span><span class="l1">1 in progress (6–8)</span><span class="l0">0 not validated (0–5)</span><span class="lx">blank = not yet covered</span></p>
<div class="cards">
<div class="card"><b>Validated</b><span>{n_val} / {len(comps)}</span><div class="bar"><i style="width:{pct}%"></i></div></div>
<div class="card"><b>Diagnostic (P1)</b><span>{esc(st['diag']) or '—'}</span></div>
<div class="card"><b>Mock 1 (P8)</b><span>{esc(st['mock1']) or '—'}</span></div>
<div class="card"><b>Mock 2 (P13)</b><span>{esc(st['mock2']) or '—'}</span></div>
<div class="card"><b>Target</b><span>{esc(st['target']) or '—'}</span></div>
</div>
{wheel_svg(st, comps)}
<p class="tip" id="tip"></p><p class="hint">Each wedge is one competence. Tap a wedge to see its name and your result.</p>
<div class="prio"><b>Next steps</b> — what to work on / retake, most SAT-heavy first:{prio_html}</div>
<table><thead><tr><th>Code</th><th>Competence</th><th>Session</th><th>Test 1</th><th>Retake</th><th>Level</th></tr></thead>
<tbody>{''.join(rows)}</tbody></table>
<footer>{esc(cfg['teacher'])} · Best level is kept after a retake · 10 questions per test · This page is private: do not share the link.</footer>
</div><script>document.querySelectorAll('.wheel g.w').forEach(g=>g.addEventListener('click',()=>{{document.getElementById('tip').textContent=g.dataset.tip}}))</script></body></html>"""
    return page

def teacher_index(students, comps, cfg, today):
    head = "".join(f"<th>{esc(c['code'])}</th>" for c in comps)
    body = []
    for st in students:
        cells = []
        for c in comps:
            lv = st["comps"][c["code"]]["level"]
            cells.append(f'<td class="lvl {"l"+str(lv) if lv is not None else "lx"}">{"" if lv is None else lv}</td>')
        n_val = sum(1 for c in comps if st["comps"][c["code"]]["level"] == 2)
        body.append(f'<tr><td><a href="docs/{esc(st["slug"])}/index.html">{esc(st["name"])}</a></td><td class="c">{n_val}</td>{"".join(cells)}</tr>')
    # per-competence class stats
    stats = []
    for c in comps:
        lv = [st["comps"][c["code"]]["level"] for st in students]
        done = [l for l in lv if l is not None]
        if done:
            stats.append(f'<td class="c">{round(100*sum(1 for l in done if l==2)/len(done))}%</td>')
        else:
            stats.append('<td class="c">—</td>')
    return f"""<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta name="robots" content="noindex,nofollow"><title>Class overview — {esc(cfg['course_title'])}</title><style>{CSS}
table{{font-size:.8rem}}td.lvl{{width:auto;padding:4px 6px}}</style></head><body><div class="wrap" style="max-width:1100px">
<h1>Class overview — {esc(cfg['course_title'])}</h1><p class="sub">{len(students)} students · updated {today} · teacher page, kept local (not published)</p>
<table><thead><tr><th>Student</th><th>Val.</th>{head}</tr></thead><tbody>{''.join(body)}
<tr class="dom"><td>% validated (of tested)</td><td></td>{''.join(stats)}</tr></tbody></table>
<footer>Student pages: {esc(cfg['base_url'])}&lt;slug&gt;/ — slugs are in passport.xlsx, column Slug.</footer></div></body></html>"""

def qr_sheet(students, cfg):
    cards = "".join(
        f'<div class="q"><img src="{esc(st["slug"])}.png" alt=""><div><b>{esc(st["name"])}</b><br>{esc(cfg["course_title"])}<br><small>{esc(cfg["base_url"])}{esc(st["slug"])}/</small></div></div>'
        for st in students)
    return f"""<!DOCTYPE html><html><head><meta charset="utf-8"><title>QR codes — passport</title><style>
body{{font-family:Helvetica,Arial,sans-serif;margin:10mm}}.grid{{display:grid;grid-template-columns:repeat(3,1fr);gap:6mm}}
.q{{border:1px dashed #999;padding:4mm;display:flex;gap:4mm;align-items:center;break-inside:avoid;font-size:10pt}}.q img{{width:28mm;height:28mm}}small{{color:#666;font-size:7pt;word-break:break-all}}
@media print{{.q{{border-color:#bbb}}}}</style></head><body><div class="grid">{cards}</div></body></html>"""

# ------------------------------------------------------------------ main
def main():
    wb = openpyxl.load_workbook(XLSX)
    cfg = read_config(wb)
    comps = read_competences(wb)
    students, changed = read_students(wb, comps)
    if changed:
        wb.save(XLSX)  # write back the generated slugs
    today = datetime.date.today().isoformat()
    SITE.mkdir(exist_ok=True); QR.mkdir(exist_ok=True)
    for st in students:
        d = SITE / st["slug"]; d.mkdir(exist_ok=True)
        (d / "index.html").write_text(student_page(st, comps, cfg, today), encoding="utf-8")
        segno.make(cfg["base_url"] + st["slug"] + "/", error="m").save(str(QR / f"{st['slug']}.png"), scale=6, border=2)
    (HERE / "overview.html").write_text(teacher_index(students, comps, cfg, today), encoding="utf-8")
    (SITE / "index.html").write_text("<!DOCTYPE html><meta charset=utf-8><meta name=robots content=noindex><title>Passport</title><p style=font-family:sans-serif>Scan your QR code to open your passport.</p>", encoding="utf-8")
    (SITE / ".nojekyll").write_text("")
    (QR / "qr-sheet.html").write_text(qr_sheet(students, cfg), encoding="utf-8")
    print(f"{len(students)} student pages written to {SITE}/  ·  QR codes in {QR}/  ·  base URL {cfg['base_url']}")
    for st in students:
        print(f"  {st['name']:<28} {cfg['base_url']}{st['slug']}/")

if __name__ == "__main__":
    main()
